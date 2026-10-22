"""Import (CSV/XLSX), merge, erasure and retention."""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone

from rapidfuzz import fuzz
from sqlalchemy import delete, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import get_settings
from app.models import (Activity, CampaignContact, Company, CompanyLocation, CompanySignal, CompanySource, Contact,
                        ContactEmail, ContactProfile, ContactSource, CrawlJob, CrawlPage, DuplicateCandidate,
                        FinanceInvoice, Note, OutreachDraft, Setting, TrainingOpportunity)
from app.models.base import (EmailDiscoveryMethod, PipelineStatus, Provenance, SourceType, VerificationStatus)
from app.services.crm import log_activity, suppress
from app.services.email_intel import classify_email, validate_syntax
from app.services.normalize import (domain_from_url, normalize_company_name, normalize_person_name, normalize_website,
                                    profile_platform, split_name)
from app.services.titles import normalize_title

MAX_BYTES = 5_000_000
MAX_ROWS = 5000
MAX_COLS = 60

FIELDS = {
    "company": ["company", "company name", "organisation", "organization", "employer", "account"],
    "website": ["website", "web", "url", "company website", "site"],
    "domain": ["domain", "company domain"],
    "contact": ["name", "full name", "contact", "contact name", "person"],
    "first_name": ["first name", "firstname", "given name"],
    "last_name": ["last name", "lastname", "surname", "family name"],
    "title": ["title", "job title", "position", "designation", "role"],
    "email": ["email", "e-mail", "email address", "work email", "business email"],
    "phone": ["phone", "telephone", "mobile", "tel", "contact number"],
    "profile": ["linkedin", "linkedin url", "profile", "profile url"],
    "city": ["city", "town"],
    "state": ["state", "region", "province"],
    "industry": ["industry", "sector"],
    "employees": ["employees", "company size", "headcount", "staff"],
    "source": ["source", "lead source"],
}


class ImportFileError(ValueError):
    pass


def parse_table(filename: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    if len(data) > MAX_BYTES:
        raise ImportFileError(f"File is larger than {MAX_BYTES // 1_000_000} MB")
    name = filename.lower()
    if name.endswith(".xlsx"):
        if not data.startswith(b"PK"):
            raise ImportFileError("Not a valid .xlsx file")
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # malformed workbook
            raise ImportFileError(f"Couldn't read the workbook: {type(e).__name__}")
        ws = wb.worksheets[0]
        rows = [["" if v is None else str(v).strip() for v in r[:MAX_COLS]]
                for r in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 1)]
        wb.close()
    elif name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        if "\x00" in text:
            raise ImportFileError("This doesn't look like a text CSV file")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [[c.strip() for c in r[:MAX_COLS]] for _, r in zip(range(MAX_ROWS + 1), csv.reader(io.StringIO(text), dialect))]
    else:
        raise ImportFileError("Upload a .csv or .xlsx file")
    rows = [r for r in rows if any(c for c in r)]
    if len(rows) < 2:
        raise ImportFileError("The file needs a header row and at least one data row")
    headers = [h or f"Column {i + 1}" for i, h in enumerate(rows[0])]
    body = [r + [""] * (len(headers) - len(r)) for r in rows[1:]]
    return headers, body


def suggest_mapping(headers: list[str]) -> dict[str, str | None]:
    out: dict[str, str | None] = {}
    used: set[str] = set()
    for f, aliases in FIELDS.items():
        best, score = None, 0
        for h in headers:
            if h in used:
                continue
            s = max(fuzz.ratio(h.lower().strip(), a) for a in aliases)
            if s > score:
                best, score = h, s
        out[f] = best if score >= 85 else None
        if out[f]:
            used.add(out[f])
    return out


@dataclass
class ImportResult:
    companies_created: int = 0
    companies_updated: int = 0
    contacts_created: int = 0
    contacts_updated: int = 0
    emails_added: int = 0
    skipped: list[dict] = field(default_factory=list)
    duplicates_flagged: int = 0
    contact_ids: list[int] = field(default_factory=list)


def _val(row: list[str], headers: list[str], mapping: dict, key: str) -> str:
    col = mapping.get(key)
    if not col or col not in headers:
        return ""
    return (row[headers.index(col)] or "").strip()[:1000]


async def _is_suppressed(db: AsyncSession, email: str | None, domain: str | None) -> bool:
    from app.models import Suppression
    conds = []
    if email:
        conds.append((Suppression.kind == "EMAIL") & (Suppression.value == email))
    if domain:
        conds.append((Suppression.kind == "DOMAIN") & (Suppression.value == domain))
    if not conds:
        return False
    from sqlalchemy import or_
    return bool(await db.scalar(select(Suppression.id).where(or_(*conds)).limit(1)))


def _fill(obj, attr: str, value) -> bool:
    """Only fills empty fields — never overwrites existing (possibly higher-confidence) data."""
    if value in (None, "") or getattr(obj, attr) not in (None, ""):
        return False
    setattr(obj, attr, value)
    return True


async def import_rows(db: AsyncSession, headers: list[str], rows: list[list[str]], mapping: dict, filename: str,
                      user_id: int | None) -> ImportResult:
    s = get_settings()
    res = ImportResult()
    src_label = f"Import: {filename}"[:120]
    company_cache: dict[str, Company] = {}
    for i, row in enumerate(rows, start=2):
        g = lambda k: _val(row, headers, mapping, k)  # noqa: E731
        company_name = g("company")
        website_raw = g("website") or g("domain")
        try:
            website = normalize_website(website_raw) if website_raw else None
        except ValueError:
            res.skipped.append({"row": i, "reason": f"Invalid website: {website_raw[:60]}"})
            continue
        domain = domain_from_url(g("domain") or website) if (website or g("domain")) else None
        email = validate_syntax(g("email")) if g("email") else None
        if g("email") and not email:
            res.skipped.append({"row": i, "reason": f"Invalid email: {g('email')[:60]}"})
            continue
        if not company_name and not domain:
            res.skipped.append({"row": i, "reason": "No company name or website"})
            continue
        if await _is_suppressed(db, email, domain or (email.rsplit("@", 1)[1] if email else None)):
            res.skipped.append({"row": i, "reason": "On the Do Not Contact list"})
            continue
        source_value = g("source") or src_label

        key = domain or normalize_company_name(company_name)
        company = company_cache.get(key)
        created_now = False
        if company is None:
            q = select(Company).where(Company.deleted_at.is_(None), Company.is_demo == s.demo_mode)
            company = (await db.execute(q.where(Company.domain == domain))).scalars().first() if domain else None
            if company is None and company_name:
                company = (await db.execute(q.where(Company.normalized_name == normalize_company_name(company_name))
                                            .order_by(Company.id))).scalars().first()
            if company is None:
                name = company_name or domain
                company = Company(legal_name=name[:300], display_name=name[:300], normalized_name=normalize_company_name(name) or name.lower(),
                                  website=website, domain=domain, domain_status="USER_PROVIDED" if domain else "UNRESOLVED",
                                  is_demo=s.demo_mode, created_by_id=user_id)
                db.add(company)
                await db.flush()
                res.companies_created += 1
                created_now = True
                log_activity(db, "created", f"Imported from {filename}", company_id=company.id, user_id=user_id)
            company_cache[key] = company
            changed = False
            for attr, value in (("industry", g("industry")[:120] or None), ("website", website)):
                changed |= _fill(company, attr, value)
            if company.domain is None and domain:
                company.domain, company.domain_status = domain, "USER_PROVIDED"
                changed = True
            emp = re.sub(r"[^\d]", "", g("employees"))
            if emp:
                changed |= _fill(company, "employee_estimate", min(int(emp), 10_000_000))
            for f, v in (("legal_name", company_name), ("website", website), ("industry", g("industry")),
                         ("employee_estimate", g("employees"))):
                if v:
                    db.add(CompanySource(company_id=company.id, field_name=f, value=str(v)[:2000], provenance=Provenance.USER_PROVIDED,
                                         source_type=SourceType.USER_IMPORT, provider=source_value[:80]))
            if (g("city") or g("state")) and not await db.scalar(select(CompanyLocation.id).where(CompanyLocation.company_id == company.id)):
                db.add(CompanyLocation(company_id=company.id, city=g("city")[:120] or None, state=g("state")[:120] or None, is_primary=True))
                changed = True
            if changed and not created_now:
                res.companies_updated += 1

        full_name = g("contact") or " ".join(x for x in (g("first_name"), g("last_name")) if x)
        if not full_name:
            continue
        norm = normalize_person_name(full_name)
        if not norm:
            res.skipped.append({"row": i, "reason": "Contact name has no usable characters"})
            continue
        existing = None
        if email:
            existing = (await db.execute(select(Contact).join(ContactEmail, ContactEmail.contact_id == Contact.id)
                                         .where(Contact.company_id == company.id, Contact.deleted_at.is_(None),
                                                ContactEmail.email == email))).scalars().first()
        if existing is None:
            existing = (await db.execute(select(Contact).where(Contact.company_id == company.id, Contact.deleted_at.is_(None),
                                                               Contact.normalized_name == norm))).scalars().first()
        if existing is None:
            similar = (await db.execute(select(Contact).where(Contact.company_id == company.id, Contact.deleted_at.is_(None)))).scalars().all()
            first, last = split_name(full_name)
            contact = Contact(company_id=company.id, full_name=full_name[:200], normalized_name=norm, first_name=first,
                              last_name=last, phone=g("phone")[:60] or None, status=PipelineStatus.REVIEW,
                              is_demo=s.demo_mode, owner_id=user_id)
            nt = normalize_title(g("title"))
            contact.original_title = g("title")[:200] or None
            contact.normalized_title, contact.department, contact.seniority = nt.normalized_title, nt.department, nt.seniority
            db.add(contact)
            await db.flush()
            res.contacts_created += 1
            for other in similar:
                score = fuzz.token_set_ratio(norm, other.normalized_name) / 100
                if score >= 0.85:
                    lo, hi = sorted((other.id, contact.id))
                    db.add(DuplicateCandidate(entity_type="contact", left_id=lo, right_id=hi, similarity=score,
                                              reasons=[f"Similar name to {other.full_name} at the same company"]))
                    res.duplicates_flagged += 1
            log_activity(db, "discovered", f"Imported from {filename} (row {i})", contact_id=contact.id, company_id=company.id, user_id=user_id)
        else:
            contact = existing
            updated = _fill(contact, "phone", g("phone")[:60] or None)
            if g("title") and not contact.original_title:
                nt = normalize_title(g("title"))
                contact.original_title = g("title")[:200]
                contact.normalized_title, contact.department, contact.seniority = nt.normalized_title, nt.department, nt.seniority
                updated = True
            res.contacts_updated += int(updated)
        res.contact_ids.append(contact.id)
        db.add(ContactSource(contact_id=contact.id, field_name="full_name", value=full_name[:2000], provenance=Provenance.USER_PROVIDED,
                             source_type=SourceType.USER_IMPORT, provider=source_value[:80]))
        if g("title"):
            db.add(ContactSource(contact_id=contact.id, field_name="title", value=g("title")[:2000], provenance=Provenance.USER_PROVIDED,
                                 source_type=SourceType.USER_IMPORT, provider=source_value[:80]))
        if email and not await db.scalar(select(ContactEmail.id).where(ContactEmail.contact_id == contact.id, ContactEmail.email == email)):
            has_primary = await db.scalar(select(ContactEmail.id).where(ContactEmail.contact_id == contact.id, ContactEmail.is_primary.is_(True)))
            db.add(ContactEmail(contact_id=contact.id, email=email, email_type=classify_email(email),
                                discovery_method=EmailDiscoveryMethod.USER_IMPORT, provenance=Provenance.USER_PROVIDED,
                                source=source_value[:120], verification_status=VerificationStatus.UNVERIFIED, is_primary=not has_primary))
            res.emails_added += 1
        profile = g("profile")
        if profile:
            try:
                url = normalize_website(profile)
                if not await db.scalar(select(ContactProfile.id).where(ContactProfile.contact_id == contact.id, ContactProfile.url == url)):
                    db.add(ContactProfile(contact_id=contact.id, platform=profile_platform(url), url=url[:1000],
                                          source_type=SourceType.USER_IMPORT, provider=source_value[:80]))
            except ValueError:
                res.skipped.append({"row": i, "reason": "Invalid profile URL (contact imported without it)"})
    await db.flush()
    return res


# --- merge -----------------------------------------------------------------------------

async def merge_companies(db: AsyncSession, keep: Company, other: Company) -> None:
    for model in (Contact, CompanySource, CompanySignal, CompanyLocation, CrawlJob, CrawlPage, Note, Activity, FinanceInvoice):
        await db.execute(update(model).where(model.company_id == other.id).values(company_id=keep.id))
    await db.execute(update(TrainingOpportunity).where(TrainingOpportunity.company_id == other.id).values(is_current=False, company_id=keep.id))
    for attr in ("website", "industry", "description", "employee_estimate", "phone", "general_email", "deal_value"):
        _fill(keep, attr, getattr(other, attr))
    if not keep.domain and other.domain:
        domain = other.domain
        other.domain = None
        await db.flush()
        keep.domain, keep.domain_status, keep.domain_confidence = domain, other.domain_status, other.domain_confidence
    other.deleted_at = datetime.now(timezone.utc)


async def merge_contacts(db: AsyncSession, keep: Contact, other: Contact) -> None:
    keep_emails = set((await db.execute(select(ContactEmail.email).where(ContactEmail.contact_id == keep.id))).scalars())
    for e in (await db.execute(select(ContactEmail).where(ContactEmail.contact_id == other.id))).scalars():
        if e.email in keep_emails:
            await db.delete(e)
        else:
            e.contact_id, e.is_primary = keep.id, False
    keep_urls = set((await db.execute(select(ContactProfile.url).where(ContactProfile.contact_id == keep.id))).scalars())
    for p in (await db.execute(select(ContactProfile).where(ContactProfile.contact_id == other.id))).scalars():
        if p.url in keep_urls:
            await db.delete(p)
        else:
            p.contact_id = keep.id
    for model in (ContactSource, Note, Activity, OutreachDraft):
        await db.execute(update(model).where(model.contact_id == other.id).values(contact_id=keep.id))
    keep_camps = set((await db.execute(select(CampaignContact.campaign_id).where(CampaignContact.contact_id == keep.id))).scalars())
    for cc in (await db.execute(select(CampaignContact).where(CampaignContact.contact_id == other.id))).scalars():
        if cc.campaign_id in keep_camps:
            await db.delete(cc)
        else:
            await db.execute(update(CampaignContact).where(CampaignContact.campaign_id == cc.campaign_id,
                                                           CampaignContact.contact_id == other.id).values(contact_id=keep.id))
    for attr in ("original_title", "normalized_title", "department", "seniority", "phone", "follow_up_at", "last_contacted_at"):
        _fill(keep, attr, getattr(other, attr))
    other.deleted_at = datetime.now(timezone.utc)


# --- erasure & retention ------------------------------------------------------------------

async def erase_contact(db: AsyncSession, contact: Contact, user_id: int | None) -> int:
    """Permanently deletes a person (data-subject erasure). Their email addresses go on
    the suppression list so discovery won't recreate them."""
    emails = list((await db.execute(select(ContactEmail.email).where(ContactEmail.contact_id == contact.id))).scalars())
    for e in emails:
        await suppress(db, "EMAIL", e, "Erasure request", user_id)
    cid = contact.id
    await db.execute(delete(DuplicateCandidate).where(DuplicateCandidate.entity_type == "contact",
                                                      (DuplicateCandidate.left_id == cid) | (DuplicateCandidate.right_id == cid)))
    await db.execute(update(FinanceInvoice).where(FinanceInvoice.contact_id == cid).values(contact_id=None))
    await db.delete(contact)  # cascades emails, profiles, sources, scores, notes, activities, drafts, memberships
    return len(emails)


async def run_retention(db: AsyncSession, dry_run: bool = True) -> dict:
    row = await db.get(Setting, "retention")
    r = (row.value if row else None) or {"activity_days": 1095, "crawl_page_days": 365, "rejected_contact_days": 180}
    now = datetime.now(timezone.utc)
    pages_q = select(CrawlPage.id).where(CrawlPage.fetched_at < now - timedelta(days=r["crawl_page_days"]))
    acts_q = select(Activity.id).where(Activity.created_at < now - timedelta(days=r["activity_days"]))
    rej_q = select(Contact).where(Contact.status == PipelineStatus.REJECTED,
                                  Contact.status_changed_at < now - timedelta(days=r["rejected_contact_days"]))
    pages = list((await db.execute(pages_q)).scalars())
    acts = list((await db.execute(acts_q)).scalars())
    rejected = list((await db.execute(rej_q)).scalars())
    if not dry_run:
        if pages:
            await db.execute(update(CompanySource).where(CompanySource.crawl_page_id.in_(pages)).values(crawl_page_id=None))
            await db.execute(update(ContactSource).where(ContactSource.crawl_page_id.in_(pages)).values(crawl_page_id=None))
            await db.execute(update(CompanySignal).where(CompanySignal.crawl_page_id.in_(pages)).values(crawl_page_id=None))
            await db.execute(delete(CrawlPage).where(CrawlPage.id.in_(pages)))
        if acts:
            await db.execute(delete(Activity).where(Activity.id.in_(acts)))
        for c in rejected:
            await db.delete(c)
    return {"settings": r, "crawl_pages": len(pages), "activities": len(acts), "rejected_contacts": len(rejected), "dry_run": dry_run}
