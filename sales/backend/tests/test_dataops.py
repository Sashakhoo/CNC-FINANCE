import io
import json
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook, load_workbook
from sqlalchemy import select, update

from app.models import (AuditLog, Company, Contact, ContactEmail, CrawlPage, DuplicateCandidate, Suppression, User)
from app.models.base import PipelineStatus, Role
from tests.conftest import finance_session

CSV = (b"Company,Website,Full Name,Job Title,Email,State,Industry,Lead Source\n"
       b"Southern Precision Sdn Bhd,southern-precision.example.my,Jane Tan,HR Manager,jane.tan@southern-precision.example.my,Johor,Automation,MyExpo 2026\n"
       b"Southern Precision Sdn Bhd,southern-precision.example.my,Daniel Lee,Engineering Manager,,Johor,Automation,MyExpo 2026\n"
       b"Bad Row Sdn Bhd,,Someone,IT Manager,not-an-email,,,\n"
       b"Blocked Co Sdn Bhd,blocked.example.my,Blocked Person,HR Manager,boss@blocked.example.my,,,\n")


async def _login(client, finance, **kw):
    client.cookies.update(finance_session(finance, **kw))
    await client.get("/api/auth/me")


async def _commit(client, data: bytes, name="leads.csv", mapping=None):
    if mapping is None:
        mapping = (await client.post("/api/import/preview", files={"file": (name, data)})).json()["mapping"]
    return await client.post("/api/import/commit", files={"file": (name, data)}, data={"mapping": json.dumps(mapping)})


async def test_preview_suggests_mapping(client, finance):
    await _login(client, finance)
    r = (await client.post("/api/import/preview", files={"file": ("leads.csv", CSV)})).json()
    assert r["total"] == 4 and r["headers"][0] == "Company"
    m = r["mapping"]
    assert (m["company"], m["website"], m["contact"], m["title"], m["email"], m["state"], m["source"]) == \
        ("Company", "Website", "Full Name", "Job Title", "Email", "State", "Lead Source")
    assert (await client.post("/api/import/preview", files={"file": ("x.exe", b"MZ")})).status_code == 422
    assert (await client.post("/api/import/preview", files={"file": ("x.xlsx", b"not a zip")})).status_code == 422


async def test_import_dedupes_fills_blanks_and_respects_suppression(client, finance, db):
    await _login(client, finance)
    await client.post("/api/suppression", json={"kind": "DOMAIN", "value": "blocked.example.my"})
    r = (await _commit(client, CSV)).json()
    assert (r["companies_created"], r["contacts_created"], r["emails_added"]) == (1, 2, 1)
    reasons = {s["reason"].split(":")[0] for s in r["skipped_rows"]}
    assert reasons == {"Invalid email", "On the Do Not Contact list"}
    co = (await db.execute(select(Company))).scalar_one()
    assert co.domain == "southern-precision.example.my" and co.industry == "Automation"

    co.industry = "Industrial Automation"          # user edit must survive a re-import
    await db.commit()
    again = (await _commit(client, CSV)).json()
    assert again["companies_created"] == 0 and again["contacts_created"] == 0
    await db.refresh(co)
    assert co.industry == "Industrial Automation"
    jane = (await db.execute(select(Contact).where(Contact.full_name == "Jane Tan"))).scalar_one()
    assert jane.department.value == "HR" and jane.status == PipelineStatus.REVIEW
    email = (await db.execute(select(ContactEmail).where(ContactEmail.contact_id == jane.id))).scalar_one()
    assert (email.discovery_method.value, email.provenance.value, email.verification_status.value, email.source) == \
        ("USER_IMPORT", "USER_PROVIDED", "UNVERIFIED", "MyExpo 2026")
    actions = [a.action for a in (await db.execute(select(AuditLog))).scalars()]
    assert actions.count("import") == 2


async def test_xlsx_formulas_are_imported_as_text_not_executed(client, finance, db):
    await _login(client, finance)
    wb = Workbook()
    ws = wb.active
    ws.append(["Company", "Name", "Title"])
    ws.append(["Formula Co Sdn Bhd", "Mei Ling", "HR Manager"])
    ws.append(["Calc Sdn Bhd", "=1+1", "IT Manager"])
    buf = io.BytesIO()
    wb.save(buf)
    r = (await _commit(client, buf.getvalue(), name="leads.xlsx")).json()
    # Formula cells are read as their saved value (never evaluated). This file was written by code, so the
    # formula has no saved value and that row has no name: no contact is created from it.
    assert r["contacts_created"] == 1
    names = set((await db.execute(select(Contact.full_name))).scalars())
    assert names == {"Mei Ling"}


async def test_export_excludes_dnc_and_neutralises_formulas(client, finance, db):
    await _login(client, finance)
    await _commit(client, CSV)
    sp = (await db.execute(select(Company.id).where(Company.domain == "southern-precision.example.my"))).scalar_one()
    evil = (await client.post("/api/contacts", json={"company_id": sp,
                                                     "full_name": "=HYPERLINK(1)"})).json()
    daniel = (await db.execute(select(Contact).where(Contact.full_name == "Daniel Lee"))).scalar_one()
    await client.post(f"/api/contacts/{daniel.id}/status", json={"status": "DO_NOT_CONTACT"})
    r = await client.get("/api/export/prospects", params={"format": "csv"})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    text = r.text
    assert "Daniel Lee" not in text and "Jane Tan" in text
    assert "'=HYPERLINK(1)" in text and evil["id"]
    x = await client.get("/api/export/prospects", params={"format": "xlsx", "department": "HR"})
    rows = list(load_workbook(io.BytesIO(x.content)).active.iter_rows(values_only=True))
    assert rows[0][0] == "Name" and sorted(r[0] for r in rows[1:]) == ["Blocked Person", "Jane Tan"]
    log = (await db.execute(select(AuditLog).where(AuditLog.action == "export"))).scalars().all()
    assert len(log) == 2 and log[1].meta["filters"] == {"department": "HR"}


async def test_viewer_cannot_export_or_import(client, finance, db):
    await _login(client, finance, username="viewer", role="admin")
    await db.execute(update(User).where(User.finance_username == "viewer").values(role=Role.VIEWER))
    await db.commit()
    assert (await client.get("/api/export/prospects")).status_code == 403
    assert (await client.post("/api/import/preview", files={"file": ("a.csv", CSV)})).status_code == 403


async def test_merge_contacts_and_dismiss(client, finance, db):
    await _login(client, finance)
    co = (await client.post("/api/companies", json={"legal_name": "Merge Test Sdn Bhd"})).json()
    a = (await client.post("/api/contacts", json={"company_id": co["id"], "full_name": "Jane Tan", "email": "jane@merge.example"})).json()
    b = (await client.post("/api/contacts", json={"company_id": co["id"], "full_name": "Jane S. Tan", "original_title": "HR Manager",
                                                  "email": "jtan@merge.example"})).json()
    db.add(DuplicateCandidate(entity_type="contact", left_id=a["id"], right_id=b["id"], similarity=0.9, reasons=["Similar name"]))
    await db.commit()
    dups = (await client.get("/api/duplicates", params={"entity_type": "contact"})).json()
    assert len(dups) == 1 and dups[0]["left"]["name"] == "Jane Tan"
    r = await client.post(f"/api/duplicates/{dups[0]['id']}/merge", json={"keep": "left"})
    assert r.status_code == 200
    kept = (await client.get(f"/api/contacts/{a['id']}")).json()
    assert {e["email"] for e in kept["emails"]} == {"jane@merge.example", "jtan@merge.example"}
    assert kept["original_title"] == "HR Manager"                       # blank filled from the merged record
    assert (await client.get(f"/api/contacts/{b['id']}")).status_code == 404
    assert (await client.post(f"/api/duplicates/{dups[0]['id']}/dismiss")).status_code == 404


async def test_erasure_removes_person_and_suppresses_email(client, finance, db):
    await _login(client, finance)
    co = (await client.post("/api/companies", json={"legal_name": "Erase Sdn Bhd"})).json()
    c = (await client.post("/api/contacts", json={"company_id": co["id"], "full_name": "Private Person",
                                                  "email": "private@erase.example"})).json()
    await client.post("/api/notes", json={"contact_id": c["id"], "body": "personal note"})
    r = (await client.delete(f"/api/contacts/{c['id']}/erase")).json()
    assert r["emails_suppressed"] == 1
    assert (await db.get(Contact, c["id"])) is None
    assert (await db.execute(select(Suppression).where(Suppression.value == "private@erase.example"))).scalar_one()
    log = (await db.execute(select(AuditLog).where(AuditLog.action == "contact_erased"))).scalar_one()
    assert "private" not in json.dumps(log.meta)


async def test_retention_preview_and_run(client, finance, db):
    await _login(client, finance)
    co = (await client.post("/api/companies", json={"legal_name": "Old Data Sdn Bhd"})).json()
    old = datetime.now(timezone.utc) - timedelta(days=400)
    db.add(CrawlPage(company_id=co["id"], url="https://old.example/", text="old", fetched_at=old))
    c = (await client.post("/api/contacts", json={"company_id": co["id"], "full_name": "Rejected Person"})).json()
    await client.post(f"/api/contacts/{c['id']}/status", json={"status": "REJECTED"})
    await db.execute(update(Contact).where(Contact.id == c["id"]).values(status_changed_at=old))
    await db.commit()
    p = (await client.get("/api/retention/preview")).json()
    assert (p["crawl_pages"], p["rejected_contacts"], p["dry_run"]) == (1, 1, True)
    r = (await client.post("/api/retention/run")).json()
    assert r["dry_run"] is False
    assert not (await db.execute(select(CrawlPage))).scalars().all()
    assert (await db.get(Contact, c["id"])) is None
